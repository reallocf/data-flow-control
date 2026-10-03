import csv
import json
import re
from pathlib import Path
from xml.etree import ElementTree as ET

import tiktoken
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN

base = Path(__file__).resolve().parent
inputs = base / "inputs"
outputs = base / "outputs"
notes = base / "notes"

budget = json.loads(
    (outputs / "section_274_budget_exact.json").read_text(
        encoding="utf-8"
    )
)

enc = tiktoken.get_encoding(budget["tokenizer"])
capacity = int(budget["legal_context_capacity"])

sections_path = max(
    inputs.glob("*sections.jsonl"),
    key=lambda p: p.stat().st_size,
)

manifest_path = next(
    outputs.glob("*recursive_prompt_manifest.jsonl")
)

xml_path = next(inputs.glob("usc*.xml"))

sections = {}

for line in sections_path.read_text(
    encoding="utf-8-sig"
).splitlines():
    if not line.strip():
        continue

    row = json.loads(line)
    ident = str(row.get("id", ""))

    if "usc_" in ident:
        key = ident.split("usc_", 1)[1]
        sections[key] = row

def main_text(text):
    marks = [
        "Editorial Notes",
        "Source Credit",
        "Statutory Notes and Related Subsidiaries",
        "Executive Documents",
    ]

    positions = [
        text.find(mark)
        for mark in marks
        if text.find(mark) > 0
    ]

    if positions:
        return text[:min(positions)]

    return text

graph = {}

with manifest_path.open(
    encoding="utf-8-sig"
) as fh:
    for line in fh:
        row = json.loads(line)
        section = str(row["source_section"])
        count = int(row["direct_refs"])
        graph[section] = [
            str(x)
            for x in row["context_sections"][:count]
        ]

seen = {"274"}
frontier = ["274"]
layers = []

while frontier:
    found = []

    for section in frontier:
        for target in graph.get(section, []):
            if target not in seen:
                seen.add(target)
                found.append(target)

    found = sorted(
        set(found),
        key=lambda x: (len(x), x),
    )

    if not found:
        break

    layers.append(found)
    frontier = found

tree = ET.parse(xml_path)
root = tree.getroot()

def local(tag):
    return tag.rsplit("}", 1)[-1]

def node_text(node):
    pieces = []

    def walk(item):
        if local(item.tag) in {
            "notes",
            "sourceCredit",
        }:
            return

        if item.text:
            pieces.append(item.text)

        for child in item:
            walk(child)

            if child.tail:
                pieces.append(child.tail)

    walk(node)

    return re.sub(
        r"\s+",
        " ",
        "".join(pieces),
    ).strip()

by_identifier = {
    node.attrib["identifier"]: node
    for node in root.iter()
    if node.attrib.get("identifier")
}

block_rows = list(
    csv.DictReader(
        (
            outputs
            / "section_274_direct_blocks_exact.csv"
        ).open(
            encoding="utf-8-sig",
            newline="",
        )
    )
)

direct_parts = []

for row in block_rows:
    ident = row["identifier"]
    section = row["section"]
    node = by_identifier[ident]

    direct_parts.append(
        sections[section]["citation"]
        + " "
        + ident
        + "\n"
        + node_text(node)
    )

context = "\n\n".join(direct_parts)

rows = []
previous_tokens = 0
last_fit = 0
first_over = None

for number, section_list in enumerate(
    layers,
    start=1,
):
    if number > 1:
        extra = "\n\n".join(
            sections[section]["citation"]
            + "\n"
            + main_text(
                sections[section]["text"]
            )
            for section in section_list
            if section in sections
        )

        if extra:
            context += "\n\n" + extra

    context_tokens = len(enc.encode(context))
    marginal = context_tokens - previous_tokens
    fits = context_tokens <= capacity

    rows.append({
        "layer": number,
        "section_count": len(section_list),
        "mapped_blocks":
            len(block_rows)
            if number == 1
            else "",
        "marginal_tokens": marginal,
        "cumulative_context_tokens":
            context_tokens,
        "legal_context_capacity":
            capacity,
        "fits": str(fits).lower(),
    })

    if fits:
        last_fit = number
    elif first_over is None:
        first_over = number

    previous_tokens = context_tokens

csv_path = (
    outputs
    / "section_274_layer_capacity.csv"
)

with csv_path.open(
    "w",
    encoding="utf-8",
    newline="",
) as fh:
    writer = csv.DictWriter(
        fh,
        fieldnames=list(rows[0]),
    )
    writer.writeheader()
    writer.writerows(rows)

note_path = (
    notes
    / "section_274_capacity_note.md"
)

table = "\n".join(
    "| "
    + str(row["layer"])
    + " | "
    + str(row["section_count"])
    + " | "
    + f'{row["marginal_tokens"]:,}'
    + " | "
    + f'{row["cumulative_context_tokens"]:,}'
    + " | "
    + row["fits"]
    + " |"
    for row in rows
)

note = f"""# Section 274 capacity trace

Legal-context capacity: {capacity:,} tokens.

| Layer | New sections | Marginal tokens | Cumulative context tokens | Fits |
| ---: | ---: | ---: | ---: | :--- |
{table}

Last complete layer within capacity: {last_fit}.
First complete layer beyond capacity: {first_over}.

## Baseline rule

Start with the target provision and the complete direct mapped context. Continue breadth-first. Before admitting another complete layer, calculate the cumulative legal-context size. Admit the layer only when the cumulative size remains within the legal-context capacity. Otherwise stop before that layer.

For section 274, this trace gives a concrete stopping point rather than an abstract capacity rule.

## Scope

Layer 1 retains the 19 structurally mapped direct blocks. Later-layer measurement follows the existing section-level recursive comparator. A later refinement can extend structural-unit mapping into those deeper layers.
"""

note_path.write_text(
    note,
    encoding="utf-8",
)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

bg = RGBColor(248, 248, 246)
ink = RGBColor(28, 28, 28)
muted = RGBColor(95, 95, 95)
line = RGBColor(75, 95, 120)
boxfill = RGBColor(232, 235, 239)

def base_slide():
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = bg
    return slide

def heading(slide, text, sub=""):
    box = slide.shapes.add_textbox(
        Inches(0.7),
        Inches(0.42),
        Inches(11.9),
        Inches(0.72),
    )
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(27)
    p.font.bold = True
    p.font.color.rgb = ink

    if sub:
        box = slide.shapes.add_textbox(
            Inches(0.72),
            Inches(1.08),
            Inches(11.8),
            Inches(0.42),
        )
        p = box.text_frame.paragraphs[0]
        p.text = sub
        p.font.size = Pt(13)
        p.font.color.rgb = muted

def text_lines(
    slide,
    values,
    x,
    y,
    w,
    h,
    size=18,
):
    box = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )

    frame = box.text_frame
    frame.clear()

    for i, value in enumerate(values):
        p = (
            frame.paragraphs[0]
            if i == 0
            else frame.add_paragraph()
        )
        p.text = value
        p.font.size = Pt(size)
        p.font.color.rgb = ink
        p.space_after = Pt(8)

slide = base_slide()

heading(
    slide,
    "Section 274: bounded dependency retrieval",
    "Concrete prompt-construction example",
)

text_lines(
    slide,
    [
        "Question: which statutory dependencies enter the legal context?",
        "Direct statutory references are mapped to structural units and deduplicated.",
        "Recursive expansion proceeds breadth-first until the next complete layer exceeds capacity.",
    ],
    0.95,
    2.15,
    11.1,
    3.2,
    22,
)

slide = base_slide()

heading(
    slide,
    "Direct context",
    "19 mapped blocks across 15 referenced sections",
)

half = (len(block_rows) + 1) // 2

text_lines(
    slide,
    [
        row["identifier"]
        for row in block_rows[:half]
    ],
    0.9,
    1.65,
    5.7,
    5.2,
    15,
)

text_lines(
    slide,
    [
        row["identifier"]
        for row in block_rows[half:]
    ],
    6.8,
    1.65,
    5.6,
    5.2,
    15,
)

slide = base_slide()

heading(
    slide,
    "Prompt capacity",
    "Tokenizer measurement",
)

text_lines(
    slide,
    [
        f'Context window: {budget["context_window"]:,} tokens',
        f'Completion reserve: {budget["output_reserve"]:,} tokens',
        f'Working margin: {budget["safety_margin"]:,} tokens',
        f'Fixed prompt text: {budget["fixed_literal_tokens"]:,} tokens',
        f'Target provision: {budget["target_provision_tokens"]:,} tokens',
        f'Legal-context capacity: {capacity:,} tokens',
        f'Direct mapped context: {budget["direct_context_tokens"]:,} tokens',
    ],
    1.0,
    1.6,
    11.0,
    4.9,
    19,
)

slide = base_slide()

heading(
    slide,
    "Layer-by-layer capacity",
    "Complete-layer admission",
)

headers = [
    "Layer",
    "New sections",
    "Marginal",
    "Cumulative",
    "Fits",
]

x = [0.9, 2.25, 4.25, 7.0, 10.35]
widths = [1.1, 1.55, 2.2, 2.7, 1.2]

for i, value in enumerate(headers):
    box = slide.shapes.add_textbox(
        Inches(x[i]),
        Inches(1.55),
        Inches(widths[i]),
        Inches(0.4),
    )
    p = box.text_frame.paragraphs[0]
    p.text = value
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = ink

for r_index, row in enumerate(rows):
    y = 2.05 + r_index * 0.55

    values = [
        str(row["layer"]),
        str(row["section_count"]),
        f'{row["marginal_tokens"]:,}',
        f'{row["cumulative_context_tokens"]:,}',
        "Yes" if row["fits"] == "true" else "No",
    ]

    for i, value in enumerate(values):
        box = slide.shapes.add_textbox(
            Inches(x[i]),
            Inches(y),
            Inches(widths[i]),
            Inches(0.4),
        )
        p = box.text_frame.paragraphs[0]
        p.text = value
        p.font.size = Pt(15)
        p.font.color.rgb = ink

slide = base_slide()

heading(
    slide,
    "Concrete stopping point",
)

labels = [
    ("§274", 0.85),
    ("Direct\nmapped\ncontext", 3.0),
    (
        f"Layer {last_fit}",
        5.45,
    ),
    (
        f"Layer {first_over}\nexceeds\ncapacity",
        8.05,
    ),
]

shapes = []

for label, x_pos in labels:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x_pos),
        Inches(2.55),
        Inches(1.75),
        Inches(1.35),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = boxfill
    shape.line.color.rgb = line

    p = shape.text_frame.paragraphs[0]
    p.text = label
    p.font.size = Pt(17)
    p.font.color.rgb = ink
    p.alignment = PP_ALIGN.CENTER

    shapes.append(shape)

for left, right in zip(
    shapes[:-1],
    shapes[1:],
):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        left.left + left.width,
        left.top + left.height // 2,
        right.left,
        right.top + right.height // 2,
    )
    connector.line.color.rgb = line

text_lines(
    slide,
    [
        f"Last complete layer within capacity: {last_fit}",
        f"First complete layer beyond capacity: {first_over}",
    ],
    3.0,
    4.75,
    7.3,
    1.4,
    20,
)

slide = base_slide()

heading(
    slide,
    "Baseline retrieval rule",
)

text_lines(
    slide,
    [
        "1  Begin with the target provision.",
        "2  Map direct statutory references to structural units.",
        "3  Remove repeated or contained targets.",
        "4  Admit the complete direct mapped context when it fits.",
        "5  Traverse later dependencies breadth-first.",
        "6  Measure the complete candidate layer before admission.",
        "7  Admit it only when cumulative context remains within capacity.",
        "8  Stop before the first complete layer beyond capacity.",
    ],
    0.95,
    1.55,
    11.2,
    5.1,
    17,
)

slide = base_slide()

heading(
    slide,
    "What remains after this example",
)

text_lines(
    slide,
    [
        "The direct structural mapping is established.",
        "The breadth-first stopping point is now measurable for section 274.",
        "The next refinement is structural-unit mapping inside deeper recursive layers.",
        "Prompt ordering and semantic ranking remain later experiments.",
    ],
    1.0,
    2.0,
    11.0,
    3.6,
    21,
)

deck_path = (
    base
    / "section_274_meeting_v3.pptx"
)

prs.save(deck_path)

summary = {
    "legal_context_capacity": capacity,
    "last_complete_layer": last_fit,
    "first_layer_beyond_capacity": first_over,
    "layers": rows,
}

summary_path = (
    outputs
    / "section_274_layer_capacity.json"
)

summary_path.write_text(
    json.dumps(
        summary,
        indent=2,
    ),
    encoding="utf-8",
)

print(
    json.dumps(
        summary,
        indent=2,
    )
)

print(deck_path)