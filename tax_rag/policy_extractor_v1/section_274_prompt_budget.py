import ast
import csv
import json
import re
from pathlib import Path
from xml.etree import ElementTree as ET

import tiktoken
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

base = Path(__file__).resolve().parent
inputs = base / "inputs"
outputs = base / "outputs"
notes = base / "notes"

model = "gpt-4.1-mini"
context_window = 1047576
output_reserve = 32768
safety_margin = 16384

try:
    enc = tiktoken.encoding_for_model(model)
except KeyError:
    enc = tiktoken.get_encoding("o200k_base")

def tokens(text):
    return len(enc.encode(text))

def norm_path(value):
    return tuple(
        x.lower()
        for x in str(value or "").split(".")
        if x
    )

def local(tag):
    return tag.rsplit("}", 1)[-1]

def node_text(node):
    parts = []

    def walk(item):
        if local(item.tag) in {"notes", "sourceCredit"}:
            return
        if item.text:
            parts.append(item.text)
        for child in item:
            walk(child)
            if child.tail:
                parts.append(child.tail)

    walk(node)
    return re.sub(r"\s+", " ", "".join(parts)).strip()

def main_cut(text):
    markers = [
        "Editorial Notes",
        "Source Credit",
        "Statutory Notes and Related Subsidiaries",
        "Executive Documents",
    ]
    hits = [
        text.find(x)
        for x in markers
        if text.find(x) > 0
    ]
    return min(hits) if hits else len(text)

def prompt_parts(path):
    tree = ast.parse(
        path.read_text(encoding="utf-8-sig")
    )

    for node in tree.body:
        if (
            isinstance(node, ast.FunctionDef)
            and node.name == "generate_policy"
        ):
            for stmt in node.body:
                if (
                    isinstance(stmt, ast.Assign)
                    and any(
                        isinstance(target, ast.Name)
                        and target.id == "prompt"
                        for target in stmt.targets
                    )
                ):
                    if not isinstance(
                        stmt.value,
                        ast.JoinedStr,
                    ):
                        raise RuntimeError(
                            "Prompt form changed."
                        )
                    return stmt.value.values

    raise RuntimeError(
        "Prompt definition not found."
    )

def section_key(node):
    if (
        isinstance(node, ast.Subscript)
        and isinstance(node.value, ast.Name)
        and node.value.id == "section"
        and isinstance(node.slice, ast.Constant)
    ):
        return node.slice.value

    return None

def render_prompt(parts, section, body, context):
    result = []

    for part in parts:
        if isinstance(part, ast.Constant):
            result.append(str(part.value))
            continue

        if not isinstance(
            part,
            ast.FormattedValue,
        ):
            raise RuntimeError(
                "Unexpected prompt element."
            )

        value = part.value

        if (
            isinstance(value, ast.Name)
            and value.id == "context"
        ):
            result.append(context)

        elif section_key(value) == "citation":
            result.append(section["citation"])

        elif (
            isinstance(value, ast.Subscript)
            and section_key(value.value) == "text"
        ):
            result.append(body)

        else:
            raise RuntimeError(
                "Unexpected prompt expression."
            )

    return "".join(result)

sections = {}

for line in (
    inputs / "title26_sections.jsonl"
).read_text(
    encoding="utf-8-sig"
).splitlines():

    if not line.strip():
        continue

    row = json.loads(line)
    ident = str(row.get("id", ""))

    if ident.startswith("26usc_"):
        sections[
            ident.removeprefix("26usc_")
        ] = row

target = sections["274"]
target_body = target["text"][
    :main_cut(target["text"])
]

candidates = set()
external = []

with (
    outputs / "rej16_reference_audit_title26.csv"
).open(
    encoding="utf-8-sig",
    newline="",
) as fh:

    for row in csv.DictReader(fh):

        if row.get("source_section") != "274":
            continue

        status = row.get("status", "")

        if status == "resolved_section":
            candidates.add((
                row.get("target_section", ""),
                norm_path(
                    row.get("target_path", "")
                ),
            ))

        elif status == "resolved_external":
            external.append(
                row.get("surface", "")
            )

selected = []

for sec, path in sorted(
    candidates,
    key=lambda x: (
        int(
            re.sub(
                r"\D",
                "",
                x[0],
            ) or 0
        ),
        x[0],
        x[1],
    ),
):
    covered = any(
        other_sec == sec
        and len(other_path) < len(path)
        and path[:len(other_path)]
            == other_path
        for other_sec, other_path
        in candidates
    )

    if not covered:
        selected.append((sec, path))

root = ET.parse(
    inputs / "usc26.xml"
).getroot()

lookup = {}

for node in root.iter():
    ident = node.attrib.get(
        "identifier",
        "",
    )

    match = re.fullmatch(
        r"/us/usc/t26/s([0-9][A-Za-z0-9-]*)(/.*)?",
        ident,
    )

    if not match:
        continue

    sec = match.group(1)
    tail = match.group(2) or ""

    path = tuple(
        x.lower()
        for x in tail.strip("/").split("/")
        if x
    )

    lookup.setdefault(
        (sec, path),
        [],
    ).append((ident, node))

blocks = []

for sec, path in selected:

    hits = lookup.get(
        (sec, path),
        [],
    )

    if len(hits) != 1:
        raise RuntimeError(
            f"Structural mapping failed for "
            f"{sec} {path}."
        )

    ident, node = hits[0]
    text = node_text(node)

    citation = sections.get(
        sec,
        {},
    ).get(
        "citation",
        sec,
    )

    rendered = (
        citation
        + " "
        + ident
        + "\n"
        + text
    )

    blocks.append({
        "section": sec,
        "path": ".".join(path),
        "identifier": ident,
        "chars": len(text),
        "tokens": tokens(rendered),
        "text": text,
        "citation": citation,
    })

if len(blocks) != 19:
    raise RuntimeError(
        f"Expected 19 direct mapped blocks, "
        f"found {len(blocks)}."
    )

direct_context = "\n\n".join(
    row["citation"]
    + " "
    + row["identifier"]
    + "\n"
    + row["text"]
    for row in blocks
)

parts = prompt_parts(
    base / "extract_v1.py"
)

fixed_text = "".join(
    str(x.value)
    for x in parts
    if isinstance(x, ast.Constant)
)

direct_prompt = render_prompt(
    parts,
    target,
    target_body,
    direct_context,
)

manifest_row = None

with (
    outputs
    / "title26_recursive_prompt_manifest.jsonl"
).open(
    encoding="utf-8-sig"
) as fh:

    for line in fh:
        row = json.loads(line)

        if str(
            row.get("source_section")
        ) == "274":
            manifest_row = row
            break

if manifest_row is None:
    raise RuntimeError(
        "Section 274 recursive row not found."
    )

full_context = "\n\n".join(
    sections[sec]["citation"]
    + "\n"
    + sections[sec]["text"][
        :main_cut(
            sections[sec]["text"]
        )
    ]
    for sec
    in manifest_row["context_sections"]
)

full_prompt = render_prompt(
    parts,
    target,
    target_body,
    full_context,
)

fixed_tokens = tokens(fixed_text)

target_tokens = tokens(
    target["citation"]
    + "\n"
    + target_body
)

direct_tokens = tokens(
    direct_context
)

direct_prompt_tokens = tokens(
    direct_prompt
)

full_prompt_tokens = tokens(
    full_prompt
)

legal_capacity = (
    context_window
    - output_reserve
    - safety_margin
    - fixed_tokens
    - target_tokens
)

remaining_after_direct = (
    legal_capacity
    - direct_tokens
)

metrics = {
    "model": model,
    "tokenizer": enc.name,
    "context_window": context_window,
    "output_reserve": output_reserve,
    "safety_margin": safety_margin,
    "fixed_literal_tokens": fixed_tokens,
    "target_provision_tokens": target_tokens,
    "direct_mapped_blocks": len(blocks),
    "direct_context_tokens": direct_tokens,
    "direct_prompt_tokens": direct_prompt_tokens,
    "legal_context_capacity":
        legal_capacity,
    "remaining_after_direct":
        remaining_after_direct,
    "external_reference_occurrences":
        len(external),
    "section_level_recursive_sections":
        int(
            manifest_row["reachable_refs"]
        ),
    "section_level_recursive_prompt_tokens":
        full_prompt_tokens,
    "direct_layer_fits":
        direct_tokens <= legal_capacity,
    "section_level_full_recursion_fits":
        full_prompt_tokens
        <= (
            context_window
            - output_reserve
            - safety_margin
        ),
}

(
    outputs
    / "section_274_budget_exact.json"
).write_text(
    json.dumps(
        metrics,
        indent=2,
    ),
    encoding="utf-8",
)

with (
    outputs
    / "section_274_direct_blocks_exact.csv"
).open(
    "w",
    encoding="utf-8",
    newline="",
) as fh:

    fields = [
        "section",
        "path",
        "identifier",
        "chars",
        "tokens",
    ]

    writer = csv.DictWriter(
        fh,
        fieldnames=fields,
    )

    writer.writeheader()

    for row in blocks:
        writer.writerow({
            key: row[key]
            for key in fields
        })

brief = f"""# Section 274 prompt-construction walkthrough

## Concrete prompt anatomy

Model: {model}
Tokenizer: {enc.name}
Model context window: {context_window:,} tokens
Output reserve: {output_reserve:,} tokens
Working safety margin: {safety_margin:,} tokens
Fixed prompt literals: {fixed_tokens:,} tokens
Target provision: {target_tokens:,} tokens
Direct statutory context: {len(blocks)} mapped blocks, {direct_tokens:,} tokens
Assembled direct prompt: {direct_prompt_tokens:,} tokens
Remaining legal-context capacity after the direct layer: {remaining_after_direct:,} tokens

The output reserve holds capacity for the completion. The safety margin is a working engineering allowance for request-format overhead and counting variation; it is not a model property and can be revised after request-level measurement.

## Direct mapped blocks

"""

for row in blocks:
    brief += (
        f"- {row['identifier']}: "
        f"{row['tokens']:,} tokens\n"
    )

brief += f"""

## Recursive comparison

The existing section-level recursion reaches {int(manifest_row['reachable_refs']):,} sections and produces an assembled prompt of {full_prompt_tokens:,} tokenizer tokens. This comparator exceeds the working input capacity, so unrestricted recursion cannot serve as the baseline for section 274.

## Baseline retrieval rule

1. Begin with the target provision.
2. Map every direct statutory reference to its structural unit.
3. Remove repeated or contained targets.
4. Admit the complete direct layer when it fits the legal-context capacity.
5. Form later layers breadth-first.
6. Admit a complete later layer only when the cumulative context remains within capacity.
7. Stop before the first complete layer that exceeds capacity.
8. Carry citation and structural identifiers with every admitted block.
9. Route external or unresolved references to additional resolution before policy construction.

## Deferred question

Prompt ordering and semantic ranking remain later experiments. They are distinct from the baseline retrieval rule.
"""

(
    notes
    / "section_274_prompt_walkthrough.md"
).write_text(
    brief,
    encoding="utf-8",
)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

bg = RGBColor(248, 248, 246)
ink = RGBColor(30, 30, 30)
muted = RGBColor(95, 95, 95)

def slide_base():
    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = bg

    return slide

def heading(slide, value, sub=""):
    box = slide.shapes.add_textbox(
        Inches(0.7),
        Inches(0.45),
        Inches(11.9),
        Inches(0.7),
    )

    p = box.text_frame.paragraphs[0]
    p.text = value
    p.font.size = Pt(27)
    p.font.bold = True
    p.font.color.rgb = ink

    if sub:
        box2 = slide.shapes.add_textbox(
            Inches(0.72),
            Inches(1.1),
            Inches(11.7),
            Inches(0.45),
        )

        p2 = box2.text_frame.paragraphs[0]
        p2.text = sub
        p2.font.size = Pt(13)
        p2.font.color.rgb = muted

def lines_box(
    slide,
    rows,
    x,
    y,
    w,
    h,
    size=18,
):
    box = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )

    tf = box.text_frame
    tf.clear()

    for i, value in enumerate(rows):
        p = (
            tf.paragraphs[0]
            if i == 0
            else tf.add_paragraph()
        )

        p.text = value
        p.font.size = Pt(size)
        p.font.color.rgb = ink
        p.space_after = Pt(7)

    return box

s = slide_base()

heading(
    s,
    "Section 274: prompt-construction walkthrough",
    "Concrete example for bounded statutory dependency retrieval",
)

lines_box(
    s,
    [
        "Question: which statutory dependencies enter the legal context?",
        "Direct references are structurally mapped and deduplicated.",
        "Later dependencies expand breadth-first until the next complete layer exceeds capacity.",
    ],
    0.95,
    2.15,
    11.1,
    3.2,
    22,
)

s = slide_base()

heading(
    s,
    "Direct statutory context",
    f"{len(blocks)} mapped blocks after containment deduplication",
)

lines_box(
    s,
    [
        row["identifier"]
        for row in blocks[:10]
    ],
    0.9,
    1.75,
    5.5,
    4.9,
    15,
)

lines_box(
    s,
    [
        row["identifier"]
        for row in blocks[10:]
    ],
    6.75,
    1.75,
    5.5,
    4.9,
    15,
)

s = slide_base()

heading(
    s,
    "Prompt anatomy",
    f"Tokenizer: {enc.name}",
)

lines_box(
    s,
    [
        f"Fixed prompt literals: {fixed_tokens:,} tokens",
        f"Target provision: {target_tokens:,} tokens",
        f"Direct statutory context: {direct_tokens:,} tokens",
        f"Assembled direct prompt: {direct_prompt_tokens:,} tokens",
        f"Output reserve: {output_reserve:,} tokens",
        f"Working safety margin: {safety_margin:,} tokens",
        f"Remaining legal-context capacity: {remaining_after_direct:,} tokens",
    ],
    1.0,
    1.6,
    11.0,
    4.9,
    19,
)

s = slide_base()

heading(
    s,
    "Why bounded expansion is necessary",
)

lines_box(
    s,
    [
        f"Direct mapped context: {direct_tokens:,} tokenizer tokens",
        f"Existing section-level recursive comparator: {int(manifest_row['reachable_refs']):,} sections",
        f"Recursive assembled prompt: {full_prompt_tokens:,} tokenizer tokens",
        f"Model context window: {context_window:,} tokens",
        "Result: the direct layer fits; unrestricted recursion does not.",
    ],
    1.0,
    1.8,
    11.0,
    4.2,
    21,
)

s = slide_base()

heading(
    s,
    "Baseline retrieval rule",
)

lines_box(
    s,
    [
        "1  Start with the target provision.",
        "2  Map direct statutory references to structural units.",
        "3  Remove repeated or contained targets.",
        "4  Admit the complete direct layer when it fits.",
        "5  Form later layers breadth-first.",
        "6  Admit a complete later layer only when cumulative context fits.",
        "7  Stop before the first complete layer that exceeds capacity.",
        "8  Route external or unresolved references to additional resolution.",
    ],
    0.95,
    1.55,
    11.2,
    5.1,
    17,
)

s = slide_base()

heading(
    s,
    "What the example establishes",
)

lines_box(
    s,
    [
        "Section 274 does not fit after unrestricted recursive expansion.",
        "The direct mapped layer fits comfortably.",
        "Breadth-first complete-layer admission gives a deterministic baseline.",
        "Prompt ordering and semantic ranking remain later experiments.",
    ],
    1.0,
    2.0,
    11.0,
    3.7,
    22,
)

deck = (
    base
    / "section_274_prompt_walkthrough_v2.pptx"
)

prs.save(deck)

print(
    json.dumps(
        metrics,
        indent=2,
    )
)

print(deck)