from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import (
    MSO_ANCHOR,
    PP_ALIGN,
)
from pptx.util import Inches, Pt


base = Path(__file__).resolve().parent

source = (
    base
    / "section_274_charles_v7.pptx"
)

output = (
    base
    / "section_274_charles_v8.pptx"
)


# Existing deck visual language.
ACCENT = RGBColor(
    0x4A,
    0x60,
    0x7A,
)

LIGHT = RGBColor(
    0xE6,
    0xEA,
    0xEF,
)

DARK = RGBColor(
    0x22,
    0x2F,
    0x3E,
)

MUTED = RGBColor(
    0x5E,
    0x6B,
    0x78,
)

WHITE = RGBColor(
    0xFF,
    0xFF,
    0xFF,
)


prs = Presentation(
    source
)

if len(prs.slides) != 7:
    raise RuntimeError(
        "Expected the v7 deck "
        f"to contain 7 slides; "
        f"found {len(prs.slides)}."
    )


# ============================================================
# Helpers
# ============================================================

def replace_existing_text(
    shape,
    text,
):
    """
    Replace text while preserving the formatting of
    the existing first run.
    """

    tf = shape.text_frame

    if not tf.paragraphs:
        raise RuntimeError(
            "Expected an existing paragraph."
        )

    paragraph = tf.paragraphs[0]

    if paragraph.runs:
        paragraph.runs[0].text = text

        for run in paragraph.runs[1:]:
            run.text = ""

    else:
        run = paragraph.add_run()
        run.text = text

    for extra in tf.paragraphs[1:]:
        extra.text = ""


def delete_shape(
    shape,
):
    element = shape._element
    element.getparent().remove(
        element
    )


def add_text(
    slide,
    text,
    left,
    top,
    width,
    height,
    *,
    size=16,
    bold=False,
    color=DARK,
    align=PP_ALIGN.LEFT,
    vertical=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(
        left,
        top,
        width,
        height,
    )

    tf = box.text_frame

    tf.clear()

    tf.margin_left = Inches(
        0.05
    )
    tf.margin_right = Inches(
        0.05
    )
    tf.margin_top = Inches(
        0.03
    )
    tf.margin_bottom = Inches(
        0.03
    )

    tf.vertical_anchor = (
        vertical
    )

    p = tf.paragraphs[0]
    p.alignment = align

    run = p.add_run()
    run.text = text

    run.font.size = Pt(
        size
    )
    run.font.bold = bold
    run.font.color.rgb = color

    return box


def add_card(
    slide,
    left,
    title,
    token_text,
    description,
):
    top = Inches(
        1.78
    )

    width = Inches(
        3.80
    )

    height = Inches(
        3.55
    )

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left,
        top,
        width,
        height,
    )

    card.fill.solid()
    card.fill.fore_color.rgb = LIGHT

    card.line.color.rgb = ACCENT
    card.line.width = Pt(
        1.25
    )

    tf = card.text_frame
    tf.clear()

    tf.margin_left = Inches(
        0.22
    )
    tf.margin_right = Inches(
        0.22
    )
    tf.margin_top = Inches(
        0.20
    )
    tf.margin_bottom = Inches(
        0.15
    )

    tf.vertical_anchor = (
        MSO_ANCHOR.MIDDLE
    )

    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER

    r = p.add_run()
    r.text = title
    r.font.size = Pt(
        18
    )
    r.font.bold = True
    r.font.color.rgb = ACCENT

    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(
        8
    )

    r = p.add_run()
    r.text = token_text
    r.font.size = Pt(
        25
    )
    r.font.bold = True
    r.font.color.rgb = DARK

    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(
        9
    )

    r = p.add_run()
    r.text = description
    r.font.size = Pt(
        14.5
    )
    r.font.color.rgb = DARK

    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(
        10
    )

    r = p.add_run()
    r.text = "3 replicates"
    r.font.size = Pt(
        13.5
    )
    r.font.bold = True
    r.font.color.rgb = MUTED

    return card


def style_cell(
    cell,
    text,
    *,
    size=14,
    bold=False,
    color=DARK,
    fill=None,
    align=PP_ALIGN.CENTER,
):
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = (
            fill
        )

    cell.margin_left = Inches(
        0.07
    )
    cell.margin_right = Inches(
        0.07
    )
    cell.margin_top = Inches(
        0.04
    )
    cell.margin_bottom = Inches(
        0.04
    )

    tf = cell.text_frame
    tf.clear()

    tf.vertical_anchor = (
        MSO_ANCHOR.MIDDLE
    )

    p = tf.paragraphs[0]
    p.alignment = align

    r = p.add_run()
    r.text = text

    r.font.size = Pt(
        size
    )
    r.font.bold = bold
    r.font.color.rgb = color


# ============================================================
# Slide 6
# Replace old "current status / conditional baseline"
# material with the actual controlled experiment.
# ============================================================

slide6 = prs.slides[5]

replace_existing_text(
    slide6.shapes[0],
    "Three-condition policy-extraction experiment",
)

replace_existing_text(
    slide6.shapes[1],
    (
        "Same Section 274 target, extraction instructions, "
        "model, and temperature; statutory context is the "
        "controlled variable"
    ),
)

# Remove old Slide-6 body.
for shape in list(
    slide6.shapes
)[2:]:
    delete_shape(
        shape
    )


card_width = Inches(
    3.80
)

gap = Inches(
    0.28
)

start = Inches(
    0.68
)


add_card(
    slide6,
    start,
    "Target-only",
    "5,926 tokens",
    (
        "Section 274 main text only"
    ),
)

add_card(
    slide6,
    start
    + card_width
    + gap,
    "Direct",
    "30,260 tokens",
    (
        "§274 + 19 direct Title-26 units\n"
        "+ 3 direct cross-code provisions"
    ),
)

add_card(
    slide6,
    start
    + (
        card_width
        + gap
    )
    * 2,
    "Bounded",
    "671,210 tokens",
    (
        "§274 + every complete dependency\n"
        "layer through Layer 3"
    ),
)


add_text(
    slide6,
    (
        "gpt-4.1-mini  ·  temperature 0  ·  "
        "same policy-extraction instructions  ·  "
        "judge disabled during generation"
    ),
    Inches(
        0.90
    ),
    Inches(
        5.62
    ),
    Inches(
        11.55
    ),
    Inches(
        0.48
    ),
    size=13.5,
    bold=False,
    color=MUTED,
    align=PP_ALIGN.CENTER,
    vertical=MSO_ANCHOR.MIDDLE,
)


add_text(
    slide6,
    (
        "Bounded context stops before complete Layer 4: "
        "the complete candidate exceeds the "
        "992,498-token statutory-context capacity."
    ),
    Inches(
        0.90
    ),
    Inches(
        6.14
    ),
    Inches(
        11.55
    ),
    Inches(
        0.62
    ),
    size=14.5,
    bold=True,
    color=ACCENT,
    align=PP_ALIGN.CENTER,
    vertical=MSO_ANCHOR.MIDDLE,
)


# ============================================================
# Slide 7
# Replace "next evaluation" with completed three-replicate
# result.
# ============================================================

slide7 = prs.slides[6]

replace_existing_text(
    slide7.shapes[0],
    "Three-replicate result",
)

# Delete old Slide-7 body.
for shape in list(
    slide7.shapes
)[1:]:
    delete_shape(
        shape
    )


add_text(
    slide7,
    (
        "Bounded retrieval broadens statutory-family discovery, "
        "but the additional context does not improve executable "
        "policy correctness"
    ),
    Inches(
        0.72
    ),
    Inches(
        1.03
    ),
    Inches(
        11.90
    ),
    Inches(
        0.52
    ),
    size=15,
    color=MUTED,
    align=PP_ALIGN.LEFT,
    vertical=MSO_ANCHOR.MIDDLE,
)


# ------------------------------------------------------------
# Result table
# ------------------------------------------------------------

table_shape = (
    slide7.shapes.add_table(
        4,
        5,
        Inches(
            0.70
        ),
        Inches(
            1.72
        ),
        Inches(
            11.93
        ),
        Inches(
            2.55
        ),
    )
)

table = table_shape.table

table.columns[0].width = (
    Inches(
        1.78
    )
)

table.columns[1].width = (
    Inches(
        2.48
    )
)

table.columns[2].width = (
    Inches(
        2.40
    )
)

table.columns[3].width = (
    Inches(
        2.72
    )
)

table.columns[4].width = (
    Inches(
        2.55
    )
)


headers = [
    "Context",
    "Family discovery\nrecall",
    "Usable-family\nrecall",
    "Exact candidate\nprecision",
    "Exact-string\nstability",
]


for col, value in enumerate(
    headers
):
    style_cell(
        table.cell(
            0,
            col,
        ),
        value,
        size=13.5,
        bold=True,
        color=WHITE,
        fill=ACCENT,
    )


result_rows = [
    (
        "Target-only",
        "28.57%",
        "23.81%",
        "6.25%",
        "0.000",
    ),
    (
        "Direct",
        "23.81%",
        "23.81%",
        "5.88%",
        "0.427",
    ),
    (
        "Bounded",
        "35.71%",
        "21.43%",
        "0.00%",
        "1.000",
    ),
]


for row_index, values in enumerate(
    result_rows,
    start=1,
):

    row_fill = (
        LIGHT
        if values[0]
        == "Bounded"
        else WHITE
    )

    for col_index, value in enumerate(
        values
    ):
        style_cell(
            table.cell(
                row_index,
                col_index,
            ),
            value,
            size=14.5,
            bold=(
                values[0]
                == "Bounded"
                or col_index
                == 0
            ),
            color=(
                ACCENT
                if (
                    values[0]
                    == "Bounded"
                    and col_index
                    in {
                        0,
                        1,
                        4,
                    }
                )
                else DARK
            ),
            fill=row_fill,
            align=(
                PP_ALIGN.LEFT
                if col_index
                == 0
                else PP_ALIGN.CENTER
            ),
        )


add_text(
    slide7,
    (
        "Exact-family recall: "
        "4.76% target-only  ·  "
        "2.38% direct  ·  "
        "0.00% bounded"
    ),
    Inches(
        0.88
    ),
    Inches(
        4.37
    ),
    Inches(
        11.55
    ),
    Inches(
        0.42
    ),
    size=12.5,
    color=MUTED,
    align=PP_ALIGN.CENTER,
    vertical=MSO_ANCHOR.MIDDLE,
)


# ------------------------------------------------------------
# Main takeaway
# ------------------------------------------------------------

callout = slide7.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(
        0.86
    ),
    Inches(
        4.89
    ),
    Inches(
        11.58
    ),
    Inches(
        1.03
    ),
)

callout.fill.solid()
callout.fill.fore_color.rgb = (
    LIGHT
)

callout.line.color.rgb = (
    ACCENT
)

callout.line.width = Pt(
    1.2
)

tf = callout.text_frame
tf.clear()

tf.margin_left = Inches(
    0.23
)
tf.margin_right = Inches(
    0.23
)
tf.margin_top = Inches(
    0.12
)
tf.margin_bottom = Inches(
    0.10
)

tf.vertical_anchor = (
    MSO_ANCHOR.MIDDLE
)

p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER

r = p.add_run()

r.text = (
    "Case-study takeaway: recursive retrieval improves "
    "legal-rule discovery breadth; faithful translation into "
    "executable DFC constraints remains the bottleneck."
)

r.font.size = Pt(
    16.5
)

r.font.bold = True
r.font.color.rgb = ACCENT


add_text(
    slide7,
    (
        "Bounded generated the same 7 unique normalized "
        "constraints in all 3 runs (mean exact-string "
        "Jaccard = 1.000), so the observed encoding errors "
        "are systematic under the frozen prompt rather than "
        "sampling noise."
    ),
    Inches(
        0.90
    ),
    Inches(
        6.07
    ),
    Inches(
        11.55
    ),
    Inches(
        0.70
    ),
    size=13.5,
    color=DARK,
    align=PP_ALIGN.CENTER,
    vertical=MSO_ANCHOR.MIDDLE,
)


# ============================================================
# Save and validate
# ============================================================

prs.save(
    output
)


check = Presentation(
    output
)

if len(
    check.slides
) != 7:
    raise RuntimeError(
        "Output slide count changed."
    )

slide6_text = "\n".join(
    shape.text
    for shape
    in check.slides[5].shapes
    if hasattr(
        shape,
        "text",
    )
)

slide7_text = "\n".join(
    shape.text
    for shape
    in check.slides[6].shapes
    if hasattr(
        shape,
        "text",
    )
)


required_slide6 = [
    "Three-condition policy-extraction experiment",
    "5,926 tokens",
    "30,260 tokens",
    "671,210 tokens",
]

required_slide7 = [
    "Three-replicate result",
    "35.71%",
    "0.00%",
    "1.000",
    "same 7 unique normalized constraints",
]


for value in required_slide6:
    if value not in slide6_text:
        raise RuntimeError(
            "Slide 6 validation failed: "
            + value
        )


for value in required_slide7:
    if value not in slide7_text:
        raise RuntimeError(
            "Slide 7 validation failed: "
            + value
        )


if (
    "Next evaluation"
    in slide7_text
):
    raise RuntimeError(
        "Old Slide-7 text remains."
    )


print(
    "PPTX_UPDATE_STATUS=PASS"
)

print(
    "SLIDES=7"
)

print(
    "SLIDE6=EXPERIMENT_DESIGN"
)

print(
    "SLIDE7=THREE_REPLICATE_RESULT"
)

print(
    f"OUTPUT={output}"
)
