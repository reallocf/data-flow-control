import hashlib
import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


base = Path(__file__).resolve().parent

source = (
    base
    / "section_274_charles_v8.pptx"
)

output = (
    base
    / "section_274_charles_v9.pptx"
)

report_path = (
    base
    / "outputs"
    / "section_274_charles_v9_validation.json"
)

report_path.parent.mkdir(
    parents=True,
    exist_ok=True,
)


DARK = RGBColor(
    0x22,
    0x2F,
    0x3E,
)


OLD_PREFIX = (
    "Bounded generated the same 7 unique normalized constraints"
)

NEW_TEXT = (
    "Bounded reproduced the same 7 unique constraints in all "
    "3 runs (Jaccard = 1.000), indicating systematic rather "
    "than sampling-driven errors."
)


def normalized(value):
    return re.sub(
        r"\s+",
        " ",
        str(value),
    ).strip()


def sha256(path):
    return hashlib.sha256(
        path.read_bytes()
    ).hexdigest()


def shape_text(shape):
    parts = []

    if getattr(
        shape,
        "has_text_frame",
        False,
    ):
        value = normalized(
            shape.text
        )

        if value:
            parts.append(
                value
            )

    if getattr(
        shape,
        "has_table",
        False,
    ):
        for row in shape.table.rows:
            for cell in row.cells:
                value = normalized(
                    cell.text
                )

                if value:
                    parts.append(
                        value
                    )

    return parts


def slide_text(slide):
    parts = []

    for shape in slide.shapes:
        parts.extend(
            shape_text(
                shape
            )
        )

    return "\n".join(
        parts
    )


v8 = Presentation(
    source
)

if len(v8.slides) != 7:
    raise RuntimeError(
        f"Expected 7 slides, found {len(v8.slides)}"
    )


# Freeze Slides 1-6 before patch.
slides_1_to_6_before = [
    slide_text(
        v8.slides[index]
    )
    for index in range(6)
]


slide7 = v8.slides[6]

matches = []

for shape in slide7.shapes:

    if not getattr(
        shape,
        "has_text_frame",
        False,
    ):
        continue

    value = normalized(
        shape.text
    )

    if OLD_PREFIX in value:
        matches.append(
            shape
        )


if len(matches) != 1:
    raise RuntimeError(
        "Expected exactly one Slide-7 stability textbox; "
        f"found {len(matches)}."
    )


shape = matches[0]

# ------------------------------------------------------------
# Replace text and explicitly force wrapping inside a safe box.
# ------------------------------------------------------------

shape.left = Inches(
    1.05
)

shape.top = Inches(
    6.05
)

shape.width = Inches(
    11.20
)

shape.height = Inches(
    0.72
)


tf = shape.text_frame
tf.clear()

tf.word_wrap = True

tf.margin_left = Inches(
    0.04
)

tf.margin_right = Inches(
    0.04
)

tf.margin_top = Inches(
    0.02
)

tf.margin_bottom = Inches(
    0.02
)

tf.vertical_anchor = (
    MSO_ANCHOR.MIDDLE
)


p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER

run = p.add_run()
run.text = NEW_TEXT

run.font.size = Pt(
    12.5
)

run.font.bold = False
run.font.color.rgb = DARK


# Save v9.
v8.save(
    output
)


# ============================================================
# Reopen and validate.
# ============================================================

v9 = Presentation(
    output
)

if len(v9.slides) != 7:
    raise RuntimeError(
        "Slide count changed."
    )


# Slides 1-6 must be textually identical.
for index in range(6):

    before = (
        slides_1_to_6_before[
            index
        ]
    )

    after = slide_text(
        v9.slides[
            index
        ]
    )

    if before != after:
        raise RuntimeError(
            f"Slide {index + 1} changed unexpectedly."
        )


slide7_text = slide_text(
    v9.slides[6]
)


if normalized(
    NEW_TEXT
) not in normalized(
    slide7_text
):
    raise RuntimeError(
        "New stability sentence not found."
    )


if OLD_PREFIX in slide7_text:
    raise RuntimeError(
        "Old clipped stability sentence remains."
    )


# ============================================================
# Preserve frozen result table exactly.
# ============================================================

tables = [
    shape.table
    for shape
    in v9.slides[6].shapes
    if getattr(
        shape,
        "has_table",
        False,
    )
]

if len(tables) != 1:
    raise RuntimeError(
        f"Expected one Slide-7 table; found {len(tables)}."
    )


table = tables[0]

values = [
    [
        cell.text.strip()
        for cell in row.cells
    ]
    for row in table.rows
]


expected = [
    [
        "Context",
        "Family discovery\nrecall",
        "Usable-family\nrecall",
        "Exact candidate\nprecision",
        "Exact-string\nstability",
    ],
    [
        "Target-only",
        "28.57%",
        "23.81%",
        "6.25%",
        "0.000",
    ],
    [
        "Direct",
        "23.81%",
        "23.81%",
        "5.88%",
        "0.427",
    ],
    [
        "Bounded",
        "35.71%",
        "21.43%",
        "0.00%",
        "1.000",
    ],
]


if values != expected:
    raise RuntimeError(
        "Slide-7 frozen result table changed."
    )


# ============================================================
# Verify stability textbox is inside slide bounds.
# ============================================================

slide_width = v9.slide_width
slide_height = v9.slide_height

target = None

for candidate in v9.slides[6].shapes:

    if not getattr(
        candidate,
        "has_text_frame",
        False,
    ):
        continue

    if normalized(
        NEW_TEXT
    ) in normalized(
        candidate.text
    ):
        target = candidate
        break


if target is None:
    raise RuntimeError(
        "Could not re-find patched textbox."
    )


if (
    target.left < 0
    or target.top < 0
    or target.left
    + target.width
    > slide_width
    or target.top
    + target.height
    > slide_height
):
    raise RuntimeError(
        "Patched textbox exceeds slide bounds."
    )


report = {
    "status":
        "PASS",

    "source":
        source.name,

    "output":
        output.name,

    "slides":
        len(
            v9.slides
        ),

    "slides_1_to_6_unchanged":
        True,

    "slide_7_table_unchanged":
        True,

    "old_stability_sentence_removed":
        True,

    "new_stability_sentence":
        NEW_TEXT,

    "textbox": {
        "left_inches":
            round(
                target.left
                / 914400,
                3,
            ),

        "top_inches":
            round(
                target.top
                / 914400,
                3,
            ),

        "width_inches":
            round(
                target.width
                / 914400,
                3,
            ),

        "height_inches":
            round(
                target.height
                / 914400,
                3,
            ),
    },

    "v8_sha256":
        sha256(
            source
        ),

    "v9_sha256":
        sha256(
            output
        ),
}


report_path.write_text(
    json.dumps(
        report,
        indent=2,
        ensure_ascii=False,
    ) + "\n",
    encoding="utf-8",
)


print(
    "V9_PATCH_STATUS=PASS"
)

print(
    "SLIDES_1_TO_6_UNCHANGED=TRUE"
)

print(
    "SLIDE7_TABLE_UNCHANGED=TRUE"
)

print(
    "OLD_CLIPPED_TEXT_REMOVED=TRUE"
)

print(
    "NEW_STABILITY_TEXT_PRESENT=TRUE"
)

print(
    "TEXTBOX_WITHIN_SLIDE=TRUE"
)

print(
    "V9_SHA256="
    + sha256(
        output
    ).upper()
)

print(
    "REPORT="
    + str(
        report_path
    )
)

print(
    "OUTPUT="
    + str(
        output
    )
)
